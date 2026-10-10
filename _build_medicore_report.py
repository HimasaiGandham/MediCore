from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(r'D:\MediCore\MediCore_Project_Progress_Report.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.72); sec.left_margin = Inches(0.82); sec.right_margin = Inches(0.82)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(45,55,72)
styles['Normal'].paragraph_format.space_after = Pt(6)
for sn, size, color in [('Title', 28, '153B5B'), ('Heading 1', 18, '153B5B'), ('Heading 2', 13, '167D9A'), ('Heading 3', 11, '153B5B')]:
    st=styles[sn]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(12 if sn!='Title' else 0); st.paragraph_format.space_after=Pt(5)

# Header and footer
header=sec.header.paragraphs[0]; header.text='MEDICORE  |  PROJECT PROGRESS'; header.style='Caption'; header.runs[0].font.color.rgb=RGBColor(95,115,130)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('MediCore project report  •  10 October 2026'); run.font.size=Pt(8); run.font.color.rgb=RGBColor(110,125,138)

p=doc.add_paragraph(style='Title'); p.add_run('MediCore Project Progress Report')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14)
r=p.add_run('Project purpose, system functions, phase status, and remaining work'); r.font.size=Pt(13); r.font.color.rgb=RGBColor(22,125,154)

p=doc.add_paragraph()
p.add_run('Current assessment. ').bold=True
p.add_run('MediCore has an implemented document-to-answer pipeline, a four-tab Streamlit application, and saved retrieval benchmarks. The strongest verified evidence is retrieval performance on the included benchmark sets. The saved authoritative answer-generation run was skipped because of API quota limits, so live LLM grounding still needs a successful run before it can be reported as verified.')

# Summary card table
t=doc.add_table(rows=2, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[Inches(2.1), Inches(2.1), Inches(2.1)]
for row in t.rows:
    for i,c in enumerate(row.cells): c.width=widths[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
labels=['PROJECT','IMPLEMENTED PHASES','GUIDELINE FILES']
vals=['MediCore','Core phases 1–7','9 authoritative JSON files']
for i in range(3):
    for cell,txt in [(t.cell(0,i),labels[i]),(t.cell(1,i),vals[i])]:
        cell.text=txt
        tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), 'EAF3F7' if cell is t.rows[0].cells[i] else 'F5F8FA'); tcPr.append(shd)
        for pp in cell.paragraphs:
            pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for rr in pp.runs:
                rr.font.size=Pt(8 if cell is t.rows[0].cells[i] else 10); rr.bold=(cell is t.rows[0].cells[i]); rr.font.color.rgb=RGBColor(22,59,91)

def bullet(text): doc.add_paragraph(text, style='List Bullet')
def phase(title, status, purpose, functions, detail):
    doc.add_heading(title, level=2)
    p=doc.add_paragraph(); p.add_run(status+' ').bold=True; p.add_run(purpose)
    p=doc.add_paragraph(); p.add_run('Main functions: ').bold=True; p.add_run(functions)
    p=doc.add_paragraph(); p.add_run('Progress: ').bold=True; p.add_run(detail)

doc.add_heading('1. Project introduction', level=1)
doc.add_paragraph('MediCore is an academic healthcare information assistant. It searches medical reference documents for passages relevant to a user question, then presents an answer with source and page or section details when available. Its intended role is to make reference information easier to find; it is not a diagnostic or treatment decision system.')

doc.add_heading('2. Main use and intended users', level=1)
doc.add_paragraph('The main use is document-grounded question answering: a person asks about a topic covered by the indexed material, MediCore retrieves matching evidence, and the interface displays the response alongside its provenance. The project is aimed at people who need to locate information in lengthy medical references, including students, researchers, and general information seekers. Clinical decisions still require qualified professional judgment.')

doc.add_heading('3. Main issue and problems encountered', level=1)
doc.add_paragraph('The user-facing problem is that useful guidance can be difficult to locate across long documents. A system must retrieve the right passage, preserve where it came from, and decline questions that the available corpus cannot support.')
doc.add_paragraph('The repository also records practical engineering constraints: Windows scientific-library thread allocation and repeated model loading motivated thread limits and model caching. Cloud answer generation can hit Gemini quota limits; the saved benchmark results explicitly record answer generation as skipped for quota. These constraints affect evaluation completeness and reliability, not the reported retrieval scores.')
doc.add_paragraph('The current knowledge catalog lists nine authoritative guideline JSON files plus a baseline reference entry. The catalog references a baseline PDF path; confirm that the PDF is present before treating baseline PDF ingestion as currently available.')

doc.add_heading('4. Major system steps', level=1)
for x in [
    'Load source documents and extract text with page or section metadata.',
    'Split document text into overlapping chunks and attach provenance fields.',
    'Create sentence-transformer embeddings for the chunks.',
    'Build or load a FAISS vector index and retrieve the most relevant chunks for a question.',
    'Generate a grounded answer with a configured cloud model or the local fallback.',
    'Display the answer, evidence, source details, guideline catalog, and search diagnostics in Streamlit.',
    'Evaluate retrieval, evidence coverage, provenance, and unsupported-question handling with benchmark cases.'
]: bullet(x)

doc.add_heading('5. Phases documented and current progress', level=1)
doc.add_paragraph('The project history and source code organize the work into seven phases. The phase labels below follow the pipeline comments in src/main.py and the project history; the user interface and evaluation suite are later project capabilities.')
phase('Phase 1 — Document ingestion','Implemented.','Extract text from PDF pages and ingest the authoritative JSON guidance records.','extract_text_from_pdf; ingest_authoritative_json; ingest_baseline_pdf; load_all_authoritative_chunks.','The ingestion layer retains source and page or section details for downstream retrieval.')
phase('Phase 2 — Chunking and metadata','Implemented.','Split source text into overlapping retrieval units and retain identifiers and provenance.','chunk_text; create_chunks_with_metadata.','Chunks carry information such as document ID, page, section, and chunk ID where the source provides it.')
phase('Phase 3 — Embeddings','Implemented.','Represent document chunks and questions as semantic vectors.','load_embedding_model; generate_embeddings_for_chunks.','The current working changes cache the loaded embedding model to reduce repeat loading.')
phase('Phase 4 — Vector index and search','Implemented.','Index vectors with FAISS and find relevant evidence for a query.','build_faiss_index; search_faiss; build_or_load_knowledge_base; search_knowledge_base.','The app and retrieval reports use the multi-source knowledge base and semantic search.')
phase('Phase 5 — Answer generation','Implemented; evaluation remains incomplete.','Construct responses from retrieved reference content using cloud or local generation paths.','build_grounded_prompt; synthesize_local_grounded_answer; generate_groq_answer; generate_answer; answer_query.','Current uncommitted changes add Gemini client reuse, model fallbacks, and explicit quota or API error messages. The saved authoritative generation results show SKIPPED (QUOTA), so live answer grounding has not been demonstrated by that run.')
phase('Phase 6 — Streamlit application','Implemented.','Provide the user-facing consultation and knowledge-base workflows.','load_rag_pipeline; process_user_query; render_sources; render_consultation_tab; render_catalog_tab; render_vector_tab; render_benchmark_tab; main.','The interface includes Clinical Consultation, Medical Guidelines Catalog, Knowledge Base & Vector Store, and Grounding Benchmark Suite tabs.')
phase('Phase 7 — Evaluation and safety checks','Implemented; ongoing refinement.','Measure retrieval quality and check grounding, citations, and refusal behavior against curated cases.','normalize_text; evaluate_retrieval; run_retrieval_evaluation; evaluate_answer; run_evaluation.','The authoritative retrieval report has 13 cases (11 supported, 2 unsupported) and records 100% top-1/top-3 accuracy, keyword recall, and provenance integrity. Baseline retrieval results cover 10 cases and record 100% for those metrics. These are small curated benchmarks, not clinical validation.')

doc.add_heading('6. Main functions by application area', level=1)
for head, body in [
('Knowledge base','load_catalog reads the central catalog; ingest_authoritative_json and ingest_baseline_pdf create source chunks; build_or_load_knowledge_base prepares or reuses the corpus and index; search_knowledge_base returns ranked evidence; get_knowledge_base_statistics supplies interface metrics.'),
('User interface','render_sidebar presents navigation and corpus context; render_consultation_tab handles question answering; render_catalog_tab browses sources; render_vector_tab exposes search and index metrics; render_benchmark_tab runs interactive benchmark cases.'),
('Evaluation','evaluate_retrieval checks expected documents, pages, keywords, and provenance; evaluate_answer checks answer grounding and refusal behavior; run_evaluation combines cases and writes machine-readable and human-readable reports.')]:
    doc.add_heading(head, level=2); doc.add_paragraph(body)

doc.add_heading('7. Progress evidence and interpretation', level=1)
doc.add_paragraph('Saved retrieval_results_authoritative.json reports 13 total cases, 11 supported cases, 2 unsupported cases, and 100% top-1/top-3 retrieval accuracy, keyword recall, and provenance integrity. retrieval_results_baseline.json reports 10 cases (8 supported, 2 unsupported) with the same retrieval metrics at 100%. These values describe the included test sets only.')
doc.add_paragraph('The saved results.json and results_authoritative.json summaries report retrieval passes, but their per-case answer status is SKIPPED (QUOTA). Treat the answer-generation portion as pending a successful evaluation run. The benchmark UI contains interactive cases, but its displayed headline rates are interface labels and should not be presented as independently verified results without a completed run.')

doc.add_heading('8. Work remaining', level=1)
for x in [
    'Run answer-generation evaluation successfully with a working quota or the local engine, and save the resulting report.',
    'Separate retrieval scores from generation and grounding scores in project documentation and UI summaries.',
    'Confirm the baseline PDF exists at the catalog path; add or correct it if baseline evaluation is expected.',
    'Refresh the README so it matches the multi-source knowledge base, current interface, benchmark behavior, and evaluation limitations.',
    'Expand benchmark coverage and test the refusal threshold and source attribution across more varied questions.',
    'Review guideline content, dates, and source links before any broader or public deployment; current evaluation is not clinical validation.'
]: bullet(x)

doc.add_heading('9. Overall status', level=1)
doc.add_paragraph('The core product path is present: documents are ingested, indexed, searched, and used by an application that displays evidence. Retrieval has strong results on the repository’s curated test sets. The clearest outstanding phase task is to complete and document answer-generation evaluation under a functioning engine, then align the README and benchmark claims with that evidence.')

doc.add_paragraph('Source basis: project source files in src/, documents/catalog.json, the nine files in documents/authoritative/, saved evaluation_results JSON reports, and recent Git history. Current working-tree changes affect src/main.py and src/evaluation.py.')
doc.save(out)
print(out)
